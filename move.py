"""Organize media files into WORKED and NOT_WORKED folders.

Run with: python move.py
"""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
import subprocess
import sys
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk


VIDEO_EXTENSIONS = {
	".mp4", ".mov", ".avi", ".mkv", ".wmv", ".flv", ".webm", ".m4v", ".mpeg", ".mpg"
}
IMAGE_EXTENSIONS = {
	".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic", ".heif"
}
SUPPORTED_EXTENSIONS = VIDEO_EXTENSIONS | IMAGE_EXTENSIONS
HEADER_VALUES = {"id", "instance id", "instance_id", "instanceid"}
INSTANCE_PATTERN = re.compile(r"^(?P<instance>.+?)(?:_(?P<duplicate>\d+))?$")
DEFAULT_EXTENSIONS = ", ".join(sorted(SUPPORTED_EXTENSIONS))
CONFIG_PATH = Path(os.environ.get("APPDATA", Path.home())) / "MediaInstanceOrganizer" / "config.json"


@dataclass(frozen=True)
class VideoPlan:
	source: Path
	filename: str
	instance_id: str
	status: str
	destination: Path


def normalize_id(value: object) -> str:
	text = str(value).strip()
	if text.endswith(".0") and text[:-2].isdigit():
		text = text[:-2]
	return text


def read_instance_ids(path: Path) -> set[str]:
	suffix = path.suffix.lower()
	values: list[object] = []

	if suffix in {".txt", ".text"}:
		values = path.read_text(encoding="utf-8-sig").splitlines()
	elif suffix == ".csv":
		with path.open("r", encoding="utf-8-sig", newline="") as file:
			values = [cell for row in csv.reader(file) for cell in row]
	elif suffix in {".xlsx", ".xlsm"}:
		try:
			from openpyxl import load_workbook
		except ImportError as error:
			raise RuntimeError("Excel support requires openpyxl. Install it with: pip install openpyxl") from error
		workbook = load_workbook(path, read_only=True, data_only=True)
		try:
			for sheet in workbook.worksheets:
				for row in sheet.iter_rows(values_only=True):
					values.extend(value for value in row if value is not None)
		finally:
			workbook.close()
	else:
		raise ValueError("Choose a TXT, CSV, XLSX, or XLSM instance list file.")

	return {
		normalized
		for value in values
		if (normalized := normalize_id(value)) and normalized.casefold() not in HEADER_VALUES
	}


def extract_instance_id(filename: str) -> str:
	match = INSTANCE_PATTERN.match(Path(filename).stem)
	return match.group("instance").strip() if match else ""


def scan_videos(
    video_folder: Path,
    output_folder: Path,
    worked_ids: set[str],
    extensions: set[str],
    scan_subfolders: bool,
) -> list[VideoPlan]:
	plans: list[VideoPlan] = []
	excluded = {path.resolve() for path in (output_folder / "WORKED", output_folder / "NOT_WORKED")}
	sources = video_folder.rglob("*") if scan_subfolders else video_folder.iterdir()
	for source in sorted(sources, key=lambda item: str(item).lower()):
		if not source.is_file() or source.suffix.lower() not in extensions:
			continue
		if any(parent.resolve() in excluded for parent in [source.parent]):
			continue
		instance_id = extract_instance_id(source.name)
		status = "WORKED" if instance_id in worked_ids else "NOT_WORKED"
		plans.append(VideoPlan(source, source.name, instance_id, status, output_folder / status))
	return plans


def open_folder(path: Path) -> None:
	path.mkdir(parents=True, exist_ok=True)
	if sys.platform == "win32":
		os.startfile(path)  # type: ignore[attr-defined]
	elif sys.platform == "darwin":
		subprocess.Popen(["open", str(path)])
	else:
		subprocess.Popen(["xdg-open", str(path)])


def open_file(path: Path) -> None:
	if sys.platform == "win32":
		os.startfile(path)  # type: ignore[attr-defined]
	elif sys.platform == "darwin":
		subprocess.Popen(["open", str(path)])
	else:
		subprocess.Popen(["xdg-open", str(path)])


class OrganizerApp:
	def __init__(self, root: tk.Tk) -> None:
		self.root = root
		self.root.title("Media Instance Organizer")
		self.root.geometry("1120x700")
		self.root.minsize(850, 500)
		self.plans: list[VideoPlan] = []
		self.worked_ids: set[str] = set()
		self.last_log: Path | None = None
		self.busy = False

		self.video_folder = tk.StringVar()
		self.list_file = tk.StringVar()
		self.output_folder = tk.StringVar()
		self.video_extensions = tk.StringVar(value=DEFAULT_EXTENSIONS)
		self.scan_subfolders = tk.BooleanVar(value=False)
		self.summary = tk.StringVar(value="Select a media folder and instance list, then click Preview / Scan.")
		self.status_text = tk.StringVar(value="Ready")
		self.load_settings()
		self.build_ui()

	def build_ui(self) -> None:
		root_frame = ttk.Frame(self.root, padding=12)
		root_frame.pack(fill="both", expand=True)
		root_frame.columnconfigure(1, weight=1)
		root_frame.rowconfigure(5, weight=1)

		self.add_path_row(root_frame, 0, "Media folder", self.video_folder, self.choose_video_folder)
		self.add_path_row(root_frame, 1, "Instance ID list", self.list_file, self.choose_list_file)
		self.add_path_row(root_frame, 2, "Output folder", self.output_folder, self.choose_output_folder)
		ttk.Label(root_frame, text="File extensions").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=4)
		ttk.Entry(root_frame, textvariable=self.video_extensions).grid(row=3, column=1, sticky="ew", pady=4)
		ttk.Checkbutton(root_frame, text="Scan subfolders", variable=self.scan_subfolders).grid(row=3, column=2, sticky="w", padx=(8, 0), pady=4)

		actions = ttk.Frame(root_frame)
		actions.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(8, 8))
		self.scan_button = ttk.Button(actions, text="Preview / Scan", command=self.preview)
		self.scan_button.pack(side="left", padx=(0, 6))
		self.copy_button = ttk.Button(actions, text="Copy Files", command=self.confirm_copy, state="disabled")
		self.copy_button.pack(side="left", padx=6)
		ttk.Button(actions, text="Open WORKED Folder", command=lambda: self.open_destination("WORKED")).pack(side="left", padx=6)
		ttk.Button(actions, text="Open NOT_WORKED Folder", command=lambda: self.open_destination("NOT_WORKED")).pack(side="left", padx=6)
		ttk.Button(actions, text="Open Log", command=self.open_log).pack(side="left", padx=6)
		ttk.Button(actions, text="Save Settings", command=self.save_settings).pack(side="left", padx=6)
		ttk.Button(actions, text="Clear / Reset", command=self.clear).pack(side="right")

		tk.Label(root_frame, textvariable=self.summary, justify="left").grid(row=5, column=0, columnspan=3, sticky="new", pady=(0, 8))

		table_frame = ttk.Frame(root_frame)
		table_frame.grid(row=6, column=0, columnspan=3, sticky="nsew")
		root_frame.rowconfigure(6, weight=1)
		columns = ("filename", "instance", "status", "destination")
		self.table = ttk.Treeview(table_frame, columns=columns, show="headings")
		headings = {"filename": "Filename", "instance": "Instance ID", "status": "Status", "destination": "Planned Destination"}
		widths = {"filename": 300, "instance": 150, "status": 130, "destination": 180}
		for column in columns:
			self.table.heading(column, text=headings[column])
			self.table.column(column, width=widths[column], anchor="w")
		y_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.table.yview)
		x_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.table.xview)
		self.table.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
		self.table.grid(row=0, column=0, sticky="nsew")
		y_scroll.grid(row=0, column=1, sticky="ns")
		x_scroll.grid(row=1, column=0, sticky="ew")
		table_frame.rowconfigure(0, weight=1)
		table_frame.columnconfigure(0, weight=1)
		ttk.Label(root_frame, textvariable=self.status_text).grid(row=7, column=0, columnspan=3, sticky="w", pady=(8, 0))

	def load_settings(self) -> None:
		try:
			settings = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
			self.video_folder.set(settings.get("video_folder", ""))
			self.list_file.set(settings.get("list_file", ""))
			self.output_folder.set(settings.get("output_folder", ""))
			self.video_extensions.set(settings.get("video_extensions", DEFAULT_EXTENSIONS))
			self.scan_subfolders.set(bool(settings.get("scan_subfolders", False)))
		except (OSError, ValueError, TypeError):
			return

	def save_settings(self, show_message: bool = True) -> None:
		try:
			CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
			settings = {
				"video_folder": self.video_folder.get().strip(),
				"list_file": self.list_file.get().strip(),
				"output_folder": self.output_folder.get().strip(),
				"video_extensions": self.video_extensions.get().strip(),
				"scan_subfolders": self.scan_subfolders.get(),
			}
			CONFIG_PATH.write_text(json.dumps(settings, indent=2), encoding="utf-8")
			if show_message:
				self.status_text.set(f"Settings saved to {CONFIG_PATH}")
		except OSError as error:
			messagebox.showerror("Settings", f"Could not save settings:\n{error}")

	def selected_extensions(self) -> set[str]:
		extensions = {value.strip().lower() for value in self.video_extensions.get().split(",") if value.strip()}
		return {value if value.startswith(".") else f".{value}" for value in extensions}

	def add_path_row(self, parent: ttk.Frame, row: int, label: str, variable: tk.StringVar, command) -> None:
		ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
		ttk.Entry(parent, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=4)
		ttk.Button(parent, text="Browse...", command=command).grid(row=row, column=2, padx=(8, 0), pady=4)

	def choose_video_folder(self) -> None:
		selected = filedialog.askdirectory(title="Select the folder containing media")
		if selected:
			self.video_folder.set(selected)
			if not self.output_folder.get():
				self.output_folder.set(selected)

	def choose_list_file(self) -> None:
		selected = filedialog.askopenfilename(
			title="Select the Instance ID list",
			filetypes=[("Supported lists", "*.txt *.csv *.xlsx *.xlsm"), ("All files", "*.*")],
		)
		if selected:
			self.list_file.set(selected)

	def choose_output_folder(self) -> None:
		selected = filedialog.askdirectory(title="Select the output / organization folder")
		if selected:
			self.output_folder.set(selected)

	def set_busy(self, value: bool) -> None:
		self.busy = value
		state = "disabled" if value else "normal"
		self.scan_button.configure(state=state)
		if value:
			self.copy_button.configure(state="disabled")

	def preview(self) -> None:
		if self.busy:
			return
		video_path = Path(self.video_folder.get().strip())
		list_path = Path(self.list_file.get().strip())
		output_path = Path(self.output_folder.get().strip() or self.video_folder.get().strip())
		extensions = self.selected_extensions()
		if not video_path.is_dir():
			messagebox.showerror("Missing media folder", "Select an existing media folder.")
			return
		if not list_path.is_file():
			messagebox.showerror("Missing instance list", "Select an existing TXT, CSV, or Excel list.")
			return
		if not output_path:
			messagebox.showerror("Missing output folder", "Select an output folder.")
			return
		if not extensions:
			messagebox.showerror("Missing file extensions", "Enter extensions separated by commas, for example: .mp4, .jpg, .png")
			return

		self.save_settings(show_message=False)
		self.set_busy(True)
		self.status_text.set("Scanning... no files will be copied.")
		threading.Thread(target=self.scan_worker, args=(video_path, list_path, output_path, extensions, self.scan_subfolders.get()), daemon=True).start()

	def scan_worker(self, video_path: Path, list_path: Path, output_path: Path, extensions: set[str], scan_subfolders: bool) -> None:
		try:
			worked_ids = read_instance_ids(list_path)
			plans = scan_videos(video_path, output_path, worked_ids, extensions, scan_subfolders)
			self.root.after(0, lambda: self.show_preview(plans, worked_ids))
		except Exception as error:
			self.root.after(0, lambda: self.scan_failed(str(error)))

	def scan_failed(self, error: str) -> None:
		self.set_busy(False)
		self.status_text.set("Scan failed.")
		messagebox.showerror("Scan failed", error)

	def show_preview(self, plans: list[VideoPlan], worked_ids: set[str]) -> None:
		self.plans = plans
		self.worked_ids = worked_ids
		self.set_busy(False)
		for item in self.table.get_children():
			self.table.delete(item)
		for plan in plans:
			self.table.insert("", "end", values=(plan.filename, plan.instance_id, plan.status, plan.destination.name))
		found_ids = {plan.instance_id for plan in plans}
		missing = worked_ids - found_ids
		worked_count = sum(plan.status == "WORKED" for plan in plans)
		self.summary.set(
			f"Total media files found: {len(plans)}    Files belonging to worked IDs: {worked_count}    "
			f"Files belonging to unworked IDs: {len(plans) - worked_count}\n"
			f"Worked Instance IDs found: {len(worked_ids & found_ids)}    "
			f"Worked Instance IDs not found in media folder: {len(missing)}"
		)
		self.status_text.set("Preview ready. Nothing has been copied.")
		self.copy_button.configure(state="normal" if plans else "disabled")

	def confirm_copy(self) -> None:
		if self.busy or not self.plans:
			return
		collisions = [plan for plan in self.plans if plan.destination.joinpath(plan.filename).exists()]
		message = f"This will copy {len(self.plans)} media file(s) into WORKED and NOT_WORKED.\n\n"
		if collisions:
			message += f"{len(collisions)} existing destination file(s) will be skipped.\n\n"
		message += "Continue? Originals will be kept and existing destination files will never be overwritten."
		if messagebox.askyesno("Confirm copy", message, icon="warning"):
			self.set_busy(True)
			self.status_text.set("Copying files...")
			threading.Thread(target=self.copy_worker, args=(list(self.plans),), daemon=True).start()

	def copy_worker(self, plans: list[VideoPlan]) -> None:
		timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
		output_folder = Path(self.output_folder.get().strip() or self.video_folder.get().strip())
		log_path = output_folder / f"media_organizer_{timestamp}.csv"
		summary_path = output_folder / f"media_organizer_{timestamp}_summary.txt"
		rows: list[dict[str, str]] = []
		copied = {"WORKED": 0, "NOT_WORKED": 0}
		skipped = 0
		errors = 0
		for plan in plans:
			destination_file = plan.destination / plan.filename
			result = "SUCCESS"
			error_message = ""
			try:
				plan.destination.mkdir(parents=True, exist_ok=True)
				if destination_file.exists():
					result = "SKIPPED"
					error_message = "Destination file already exists"
					skipped += 1
				else:
					shutil.copy2(str(plan.source), str(destination_file))
					copied[plan.status] += 1
			except Exception as error:
				result = "ERROR"
				error_message = str(error)
				errors += 1
			rows.append({
				"timestamp": datetime.now().isoformat(timespec="seconds"),
				"original_path": str(plan.source),
				"filename": plan.filename,
				"instance_id": plan.instance_id,
				"status": plan.status,
				"destination": str(destination_file),
				"copy_result": result,
				"error": error_message,
			})

		try:
			output_folder.mkdir(parents=True, exist_ok=True)
			with log_path.open("w", encoding="utf-8", newline="") as file:
				writer = csv.DictWriter(file, fieldnames=list(rows[0]) if rows else ["timestamp"])
				writer.writeheader()
				writer.writerows(rows)
			found_ids = {plan.instance_id for plan in plans}
			missing = sorted(self.worked_ids - found_ids)
			summary_path.write_text(
				f"Media organizer summary\n"
				f"Completed: {datetime.now().isoformat(timespec='seconds')}\n"
				f"Copied to WORKED: {copied['WORKED']}\n"
				f"Copied to NOT_WORKED: {copied['NOT_WORKED']}\n"
				f"Skipped: {skipped}\nErrors: {errors}\n"
				f"Worked IDs not found: {', '.join(missing) if missing else 'None'}\n",
				encoding="utf-8",
			)
			self.root.after(0, lambda: self.copy_finished(log_path, copied, skipped, errors))
		except Exception as error:
			self.root.after(0, lambda: self.scan_failed(f"Files were processed, but the log could not be written: {error}"))

	def copy_finished(self, log_path: Path, copied: dict[str, int], skipped: int, errors: int) -> None:
		self.set_busy(False)
		self.last_log = log_path
		self.plans = []
		self.copy_button.configure(state="disabled")
		self.status_text.set(f"Copy complete. Log: {log_path.name}")
		messagebox.showinfo(
			"Copy complete",
			f"WORKED: {copied['WORKED']}\nNOT_WORKED: {copied['NOT_WORKED']}\nSkipped: {skipped}\nErrors: {errors}\n\nOriginal files were kept.\nLog: {log_path}",
		)

	def open_destination(self, name: str) -> None:
		base = Path(self.output_folder.get().strip() or self.video_folder.get().strip())
		if base:
			open_folder(base / name)

	def open_log(self) -> None:
		if not self.last_log or not self.last_log.exists():
			messagebox.showinfo("Open log", "No log has been created yet.")
			return
		open_file(self.last_log)

	def clear(self) -> None:
		if self.busy:
			return
		self.video_folder.set("")
		self.list_file.set("")
		self.output_folder.set("")
		self.video_extensions.set(DEFAULT_EXTENSIONS)
		self.scan_subfolders.set(False)
		self.plans = []
		self.worked_ids = set()
		self.last_log = None
		for item in self.table.get_children():
			self.table.delete(item)
		self.summary.set("Select a media folder and instance list, then click Preview / Scan.")
		self.status_text.set("Ready")
		self.copy_button.configure(state="disabled")


if __name__ == "__main__":
	application = tk.Tk()
	OrganizerApp(application)
	application.mainloop()
